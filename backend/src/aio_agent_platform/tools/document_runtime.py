"""Standalone document worker, shipped as source to the sandbox per invocation.

No platform imports: untrusted parsers, OCR and Office conversion run as the
sandbox user. Stdout is a JSON envelope; image bytes are a separate transient field.
"""

from __future__ import annotations

import base64
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from uuid import uuid4

MAX_BYTES = 50 * 1024 * 1024
MAX_PIXELS = 25_000_000
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
OFFICE_EXTS = {".docx", ".xlsx", ".pptx", ".doc", ".xls", ".ppt"}


def integer(args, key, default, low, high):
    value = args.get(key, default)
    if type(value) is not int or not low <= value <= high:
        raise ValueError(f"{key} must be an integer between {low} and {high}")
    return value


def source_path(root, value):
    if not isinstance(value, str) or not value or "\x00" in value:
        raise ValueError("path must be a nonempty workspace file path")
    root = Path(root).resolve(strict=True)
    path = (root / value).resolve(strict=True)
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError("Path is outside the current workspace or is not a file")
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("File exceeds the 50 MiB document limit")
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("Expanded Office document exceeds 200 MiB")
    return root, path


def command(argv, timeout=60):
    try:
        result = subprocess.run(argv, capture_output=True, timeout=timeout, check=False)
    except FileNotFoundError:
        raise ValueError(f"Missing document dependency: {argv[0]}. Rebuild the sandbox image.")
    except subprocess.TimeoutExpired:
        raise ValueError(f"Document operation timed out: {Path(argv[0]).name}")
    if result.returncode:
        raise ValueError(f"{Path(argv[0]).name} failed: {result.stderr.decode(errors='replace')[-1200:]}")
    return result.stdout


def office_convert(path, directory, extension):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise ValueError(
            "Office layout rendering/legacy conversion is optional and not installed. "
            "Use native read_document for DOCX/XLSX/PPTX, or upload a PDF to render. "
            "Administrators can build the sandbox with INSTALL_OFFICE_RENDERER=true."
        )
    # A unique profile avoids contention between users and running Office instances.
    profile = directory / "profile"
    command([
        executable, f"-env:UserInstallation={profile.as_uri()}", "--headless",
        "--convert-to", extension, "--outdir", str(directory), str(path),
    ], timeout=90)
    converted = directory / f"{path.stem}.{extension}"
    if not converted.is_file() or converted.stat().st_size == 0:
        raise ValueError("LibreOffice did not produce a converted document")
    if converted.stat().st_size > MAX_BYTES:
        raise ValueError("Converted document exceeds 50 MiB")
    return converted


def page_range(args, count, maximum=20, default=10):
    start = integer(args, "start_page", 1, 1, max(count, 1))
    end = integer(args, "end_page", min(count, start + default - 1), start, count)
    if end - start + 1 > maximum:
        raise ValueError(f"Read at most {maximum} pages per call")
    return start, end


def raster_page(pdf, index, dpi):
    page = pdf[index]
    try:
        width, height = page.get_size()
        if width * height * (dpi / 72) ** 2 > MAX_PIXELS:
            raise ValueError("Rendered page exceeds 25 million pixels; use a lower dpi")
        bitmap = page.render(scale=dpi / 72)
        try:
            return bitmap.to_pil().copy()
        finally:
            bitmap.close()
    finally:
        page.close()


def ocr_image(image, language, directory, layout="auto"):
    filename = directory / "ocr.png"
    image.save(filename, "PNG")
    mode = {"auto": "3", "block": "6", "sparse": "11"}[layout]
    text = command(["tesseract", str(filename), "stdout", "-l", language, "--psm", mode], timeout=45)
    if not text.strip() and layout == "auto":
        text = command(["tesseract", str(filename), "stdout", "-l", language, "--psm", "6"], timeout=45)
    return text.decode("utf-8", errors="replace").strip()


def image_open(path):
    from PIL import Image, ImageOps

    Image.MAX_IMAGE_PIXELS = MAX_PIXELS
    with Image.open(path) as original:
        if original.width * original.height > MAX_PIXELS:
            raise ValueError("Image exceeds 25 million pixels")
        oriented = ImageOps.exif_transpose(original)
        rgba = oriented.convert("RGBA")
        background = Image.new("RGBA", rgba.size, "white")
        image = Image.alpha_composite(background, rgba).convert("RGB")
        return image, getattr(original, "n_frames", 1)


def add_text(blocks, text, location, kind="text", **extra):
    # Long paragraphs remain readable through ordinary block pagination.
    if text:
        for offset in range(0, len(text), 6000):
            blocks.append({
                "kind": kind, "location": {**location, "char_offset": offset},
                "text": text[offset:offset + 6000], **extra,
            })


def add_row(blocks, cells, location):
    # Preserve column boundaries without allowing one giant row to bypass paging.
    for col, value in enumerate(cells, 1):
        if value is not None and str(value):
            add_text(blocks, str(value), {**location, "column": col}, kind="table_cell")


def read_pdf(path, args, directory):
    import pdfplumber
    import pypdfium2 as pdfium

    blocks, warnings = [], []
    with pdfplumber.open(path) as pdf:
        start, end = page_range(args, len(pdf.pages))
        for number in range(start, end + 1):
            page = pdf.pages[number - 1]
            text = page.extract_text() or ""
            mode = args.get("ocr", "auto")
            if mode == "always" or (mode == "auto" and not text.strip()):
                with pdfium.PdfDocument(str(path)) as raster:
                    image = raster_page(raster, number - 1, 144)
                text = ocr_image(image, args.get("ocr_language", "chi_sim+eng"), directory,
                                 args.get("ocr_layout", "auto"))
                add_text(blocks, text, {"page": number}, "ocr_text")
                warnings.append(f"Page {number}: OCR text; table structure and recognition may be imperfect.")
            else:
                add_text(blocks, text, {"page": number})
                for index, table in enumerate(page.find_tables(), 1):
                    for row, values in enumerate(table.extract(), 1):
                        add_row(blocks, values, {
                            "page": number, "table": index, "row": row,
                            "bbox": list(table.bbox), "bbox_unit": "pt",
                        })
            if not text.strip():
                warnings.append(f"Page {number}: no text extracted; use render_document/view_image or enable OCR.")
            page.close()
        return blocks, {
            "total_pages": len(pdf.pages), "start_page": start, "end_page": end,
            "next_page": end + 1 if end < len(pdf.pages) else None,
        }, warnings


def read_docx(path, args):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    blocks, headings, stack = [], [], []
    paragraph_number = table_number = 0
    requested = args.get("section")
    active_level = None
    found = requested is None
    active = requested is None
    for element in document.element.body:
        if element.tag.endswith("}p"):
            paragraph_number += 1
            paragraph = Paragraph(element, document)
            style = paragraph.style.name if paragraph.style else ""
            if style.startswith("Heading ") and style[8:].isdigit():
                level = int(style[8:])
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, paragraph.text))
                headings.append(paragraph.text)
                if requested is not None:
                    if paragraph.text == requested:
                        active, found, active_level = True, True, level
                    elif active and active_level is not None and level <= active_level:
                        active = False
            if active:
                add_text(blocks, paragraph.text, {
                    "paragraph": paragraph_number, "section": " / ".join(v for _, v in stack),
                })
        elif element.tag.endswith("}tbl"):
            table_number += 1
            if active:
                table = Table(element, document)
                for row, values in enumerate(table.rows, 1):
                    add_row(blocks, [cell.text for cell in values.cells], {
                        "table": table_number, "row": row,
                        "section": " / ".join(v for _, v in stack),
                    })
    if not found:
        raise ValueError(f"Section not found: {requested}. Available headings: {headings[:40]}")
    return blocks, {"section": requested, "headings": headings[:200]}, [
        "DOCX locations are paragraphs/headings, not page numbers. Use page selection or render_document for LibreOffice pagination."
    ]


def read_xlsx(path, args):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet_name = args.get("sheet", workbook.sheetnames[0])
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Unknown sheet: {sheet_name}. Available: {workbook.sheetnames}")
        sheet = workbook[sheet_name]
        selection = args.get("cell_range") or f"A1:{get_column_letter(min(sheet.max_column or 1, 50))}{min(sheet.max_row or 1, 200)}"
        min_col, min_row, max_col, max_row = range_boundaries(selection)
        if (not all(type(n) is int for n in (min_col, min_row, max_col, max_row))
                or not 1 <= min_col <= max_col <= 16384
                or not 1 <= min_row <= max_row <= 1048576
                or (max_col - min_col + 1) * (max_row - min_row + 1) > 20000):
            raise ValueError("cell_range must be an A1 rectangle of at most 20,000 cells")
        blocks = []
        for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if cell.value is not None:
                    add_text(blocks, str(cell.value), {"sheet": sheet_name, "cell": cell.coordinate},
                             "formula" if cell.data_type == "f" else "table_cell")
        return blocks, {
            "sheet": sheet_name, "sheets": workbook.sheetnames, "cell_range": selection,
            "sheet_rows": sheet.max_row, "sheet_columns": sheet.max_column,
        }, ["Formulas are returned as expressions; this tool does not calculate them. Only the selected sheet/range was read."]
    finally:
        workbook.close()


def read_pptx(path, args):
    from pptx import Presentation

    presentation = Presentation(path)
    start, end = page_range(args, len(presentation.slides))
    blocks = []

    def shapes(items, slide, prefix=""):
        for index, shape in enumerate(items, 1):
            shape_id = f"{prefix}{index}"
            location = {"slide": slide, "shape": shape_id}
            if shape.has_text_frame:
                for paragraph, value in enumerate(shape.text_frame.paragraphs, 1):
                    add_text(blocks, value.text, {**location, "paragraph": paragraph})
            if shape.has_table:
                for row, values in enumerate(shape.table.rows, 1):
                    add_row(blocks, [cell.text for cell in values.cells], {**location, "row": row})
            if hasattr(shape, "shapes"):
                shapes(shape.shapes, slide, shape_id + ".")

    for number in range(start, end + 1):
        shapes(presentation.slides[number - 1].shapes, number)
    return blocks, {
        "total_pages": len(presentation.slides), "start_page": start, "end_page": end,
        "next_page": end + 1 if end < len(presentation.slides) else None,
    }, ["Embedded images/charts are not OCRed in native PPTX extraction; render slides and view_image to inspect them."]


def extract(path, args, directory):
    ext = path.suffix.lower()
    warnings = []
    if ext in {".doc", ".xls", ".ppt"}:
        ext = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}[ext]
        path = office_convert(path, directory, ext[1:])
        warnings.append("Legacy Office document converted with LibreOffice.")
    if ext == ".docx" and ("start_page" in args or "end_page" in args):
        if args.get("section"):
            raise ValueError("section and page selection cannot be combined")
        path = office_convert(path, directory, "pdf")
        ext = ".pdf"
        warnings.append("Page numbers follow LibreOffice pagination, which may differ from Microsoft Word.")
    if ext == ".pdf":
        result = read_pdf(path, args, directory)
    elif ext == ".docx":
        result = read_docx(path, args)
    elif ext == ".xlsx":
        result = read_xlsx(path, args)
    elif ext == ".pptx":
        result = read_pptx(path, args)
    elif ext in IMAGE_EXTS:
        if args.get("ocr") == "never":
            raise ValueError("Reading an image as text requires OCR; use view_image for visual inspection")
        image, frames = image_open(path)
        blocks = []
        add_text(blocks, ocr_image(image, args.get("ocr_language", "chi_sim+eng"), directory,
                                  args.get("ocr_layout", "auto")),
                 {"frame": 1}, "ocr_text")
        result = blocks, {"frame": 1, "total_frames": frames}, ["OCR text may contain recognition errors; only the first frame was read."]
    else:
        raise ValueError(f"Unsupported document format: {ext}")
    return result[0], result[1], warnings + result[2]


def render(path, args, root, directory):
    import pypdfium2 as pdfium

    converted = path.suffix.lower() in OFFICE_EXTS
    if converted:
        path = office_convert(path, directory, "pdf")
    elif path.suffix.lower() != ".pdf":
        raise ValueError("render_document supports PDF and Office files; use view_image for images")
    dpi = integer(args, "dpi", 120, 72, 144)
    # Resolve before writing to reject a symlinked output directory escaping the workspace.
    parent = root / "document-renders"
    if parent.is_symlink() or not parent.resolve().is_relative_to(root):
        raise ValueError("Unsafe document-renders directory")
    parent.mkdir(exist_ok=True)
    destination = parent / uuid4().hex
    destination.mkdir()
    files = []
    try:
        with pdfium.PdfDocument(str(path)) as pdf:
            start, end = page_range(args, len(pdf), maximum=5, default=1)
            count = len(pdf)
            for number in range(start, end + 1):
                image = raster_page(pdf, number - 1, dpi)
                target = destination / f"page-{number}.png"
                image.save(target, "PNG")
                files.append({"path": target.relative_to(root).as_posix(), "page": number,
                              "width": image.width, "height": image.height})
    except Exception:
        shutil.rmtree(destination)
        raise
    return {"files": files, "total_pages": count, "start_page": start, "end_page": end,
            "next_page": end + 1 if end < count else None,
            "warnings": ["Office pages use LibreOffice pagination and installed fonts."] if converted else []}


def run(operation, args, root):
    root, path = source_path(root, args.get("path"))
    source = path.relative_to(root).as_posix()
    if args.get("ocr", "auto") not in {"auto", "always", "never"}:
        raise ValueError("ocr must be auto, always or never")
    if args.get("ocr_language", "chi_sim+eng") not in {"chi_sim+eng", "eng", "chi_sim"}:
        raise ValueError("Unsupported OCR language")
    if args.get("ocr_layout", "auto") not in {"auto", "block", "sparse"}:
        raise ValueError("Unsupported OCR layout")
    with tempfile.TemporaryDirectory(prefix="aio-document-") as temporary:
        directory = Path(temporary)
        if operation == "view_image":
            if path.suffix.lower() not in IMAGE_EXTS:
                raise ValueError("Unsupported image format; render documents to PNG first")
            image, frames = image_open(path)
            original_size = list(image.size)
            image.thumbnail((2048, 2048))
            output = io.BytesIO()
            image.save(output, "JPEG", quality=85)
            return {"result": {"source": source, "original_size": original_size,
                               "display_size": list(image.size), "frame": 1, "total_frames": frames},
                    "image": "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode()}
        if operation == "render_document":
            return {"result": {"source": source, **render(path, args, root, directory)}}
        if operation not in {"read_document", "document_search"}:
            raise ValueError("Unknown document operation")
        # Reject selectors that would otherwise silently search the wrong content.
        ext = path.suffix.lower()
        if any(k in args for k in ("sheet", "cell_range")) and ext not in {".xlsx", ".xls"}:
            raise ValueError("sheet/cell_range only apply to Excel files")
        if "section" in args and ext not in {".docx", ".doc"}:
            raise ValueError("section only applies to Word files")
        if any(k in args for k in ("start_page", "end_page")) and ext in IMAGE_EXTS | {".xlsx", ".xls"}:
            raise ValueError("For images use frame 1; for Excel reading use sheet/cell_range, not page numbers")
        offset = integer(args, "offset", 0, 0, 10_000_000)
        limit = integer(args, "limit", 20 if operation == "document_search" else 40, 1, 100)
        query = args.get("query", "")
        if operation == "document_search" and (not isinstance(query, str) or not 1 <= len(query.strip()) <= 500):
            raise ValueError("query must contain 1-500 nonblank characters")
        blocks, coverage, warnings = extract(path, args, directory)
        selected, budget, cursor = [], 0, offset
        while cursor < len(blocks):
            block = blocks[cursor]
            if operation == "document_search":
                searchable = block["text"]
                # Search across internal chunk boundaries in the same source paragraph/page.
                if cursor + 1 < len(blocks):
                    following = blocks[cursor + 1]
                    location = {k: v for k, v in block["location"].items() if k != "char_offset"}
                    next_location = {k: v for k, v in following["location"].items() if k != "char_offset"}
                    if (location == next_location and following["kind"] == block["kind"]
                            and following["location"]["char_offset"] == block["location"]["char_offset"] + len(searchable)):
                        searchable += following["text"][:len(query) - 1]
                match = re.search(re.escape(query), searchable, re.IGNORECASE)
                if match is None or match.start() >= len(block["text"]):
                    cursor += 1
                    continue
                begin = max(0, match.start() - 120)
                block = {**block, "text": searchable[begin:match.end() + 240],
                         "excerpt_offset": begin}
            cost = len(json.dumps(block, ensure_ascii=False))
            if selected and (len(selected) >= limit or budget + cost > 32000):
                break
            selected.append({"block_index": cursor, **block})
            budget += cost
            cursor += 1
        return {"result": {
            "source": source, "format": path.suffix.lower().lstrip("."),
            "coverage": coverage, "blocks": selected, "offset": offset,
            "total_blocks_in_selection": len(blocks),
            "next_offset": cursor if cursor < len(blocks) else None, "warnings": warnings,
        }}


if __name__ == "__main__":
    try:
        # Bound the worker independently of the parent tool timeout.
        import signal

        def deadline(_signum, _frame):
            raise TimeoutError("Document operation exceeded its 170-second time budget")

        signal.signal(signal.SIGALRM, deadline)
        signal.alarm(170)
        request = json.loads(base64.b64decode(sys.argv[1]))
        print(json.dumps(run(request["operation"], request["args"], request["root"]), ensure_ascii=False))
    except ModuleNotFoundError as error:
        print(json.dumps({"error": f"Missing document dependency: {error.name}. Rebuild the sandbox image."}))
        sys.exit(1)
    except Exception as error:
        print(json.dumps({"error": f"{type(error).__name__}: {error}"}, ensure_ascii=False))
        sys.exit(1)
