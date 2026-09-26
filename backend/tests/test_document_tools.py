"""Real document fixtures and model/executor contracts; no live DB or Docker needed."""

import asyncio
import base64
import io
import json
import os
import shlex
import shutil
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from aio_agent_platform.core.agent import AgentLoop
from aio_agent_platform.core.chat_history import ChatTurnRecorder
from aio_agent_platform.core.task_scope import call_key, completed_calls
from aio_agent_platform.llm.client import LLMChunk, ToolCall, sanitize_messages_for_trace
from aio_agent_platform.tools.document_runtime import run
from aio_agent_platform.tools.document_tools import DOCUMENT_TOOLS
from aio_agent_platform.tools.executor import ToolExecutor
from aio_agent_platform.tools.registry import ToolRegistry


@pytest.fixture
def documents(tmp_path):
    doc = Document()
    doc.add_heading("Overview", 1)
    doc.add_paragraph("Original contract overview.")
    doc.add_heading("Budget", 1)
    doc.add_paragraph("Approved budget: 1200.")
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Item", "Amount"), ("Laptop", "1200")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    doc.add_heading("Details", 2)
    doc.add_paragraph("Budget details remain in section.")
    doc.add_heading("Appendix", 1)
    doc.add_paragraph("Excluded appendix.")
    doc.save(tmp_path / "contract.docx")

    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["Unused", 10])
    sheet = wb.create_sheet("Budget")
    sheet.append(["Item", "Amount"])
    sheet.append(["Laptop", 1200])
    sheet.append(["Total", "=SUM(B2:B2)"])
    wb.save(tmp_path / "budget.xlsx")

    deck = Presentation()
    for content in ["First slide", "Budget approved"]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = content
    deck.save(tmp_path / "slides.pptx")

    pdf = canvas.Canvas(str(tmp_path / "report.pdf"), pagesize=(400, 400))
    for number in range(1, 13):
        pdf.drawString(40, 350, f"Budget page {number}: approved 1200")
        if number == 1:
            for x in [40, 140, 240]:
                pdf.line(x, 200, x, 280)
            for y in [200, 240, 280]:
                pdf.line(40, y, 240, y)
            pdf.drawString(50, 255, "Item")
            pdf.drawString(150, 255, "Amount")
            pdf.drawString(50, 215, "Laptop")
            pdf.drawString(150, 215, "1200")
        pdf.showPage()
    pdf.save()
    Image.new("RGB", (2400, 800), "navy").save(tmp_path / "chart.png")
    return tmp_path


def call(root, tool="read_document", **args):
    return run(tool, args, root)["result"]


def test_docx_sections_tables_and_pagination(documents):
    result = call(documents, path="contract.docx", section="Budget", limit=2)
    blocks = list(result["blocks"])
    while result["next_offset"] is not None:
        result = call(documents, path="contract.docx", section="Budget", limit=2,
                      offset=result["next_offset"])
        blocks.extend(result["blocks"])
    texts = " ".join(b["text"] for b in blocks)
    assert "1200" in texts and "Budget details" in texts
    assert "Excluded" not in texts and "overview" not in texts
    cell = next(b for b in blocks if b["text"] == "Laptop")
    assert cell["location"]["table"] == 1 and cell["location"]["row"] == 2
    assert len({b["block_index"] for b in blocks}) == len(blocks)


def test_xlsx_sheet_cell_range_and_formulas(documents):
    result = call(documents, path="budget.xlsx", sheet="Budget", cell_range="B2:B3")
    assert [(b["location"]["cell"], b["text"]) for b in result["blocks"]] == [
        ("B2", "1200"), ("B3", "=SUM(B2:B2)"),
    ]
    assert result["blocks"][1]["kind"] == "formula"
    assert result["coverage"]["sheets"] == ["Summary", "Budget"]
    with pytest.raises(ValueError, match="20,000"):
        call(documents, path="budget.xlsx", cell_range="A1:ZZ10000")


def test_pdf_pages_tables_search_and_coverage(documents):
    result = call(documents, path="report.pdf", start_page=1, end_page=1)
    assert result["coverage"]["next_page"] == 2
    cell = next(b for b in result["blocks"] if b["kind"] == "table_cell" and b["text"] == "1200")
    assert cell["location"]["page"] == 1 and cell["location"]["bbox"]
    result = call(documents, "document_search", path="report.pdf", query="APPROVED", limit=1)
    assert result["coverage"]["end_page"] == 10
    assert len(result["blocks"]) == 1 and result["next_offset"] is not None
    continuation = call(documents, "document_search", path="report.pdf", query="APPROVED",
                        offset=result["next_offset"], limit=1)
    assert continuation["blocks"][0]["location"]["page"] == 2


def test_pptx_slide_selection(documents):
    result = call(documents, path="slides.pptx", start_page=2)
    assert result["blocks"][0]["text"] == "Budget approved"
    assert result["blocks"][0]["location"] == {"slide": 2, "shape": "1", "paragraph": 1, "char_offset": 0}


def test_lightweight_mode_reads_office_without_libreoffice(documents, monkeypatch):
    monkeypatch.setattr("aio_agent_platform.tools.document_runtime.shutil.which", lambda _: None)
    for filename in ["contract.docx", "budget.xlsx", "slides.pptx"]:
        assert call(documents, path=filename)["blocks"]
        with pytest.raises(ValueError, match="optional and not installed"):
            call(documents, "render_document", path=filename)
    assert call(documents, "render_document", path="report.pdf")["files"]


def test_search_matches_across_internal_text_chunks(documents):
    doc = Document()
    doc.add_paragraph("x" * 5998 + "Budget approved" + "z" * 100)
    doc.save(documents / "long.docx")
    result = call(documents, "document_search", path="long.docx", query="budget approved")
    assert len(result["blocks"]) == 1
    assert "Budget approved" in result["blocks"][0]["text"]


def test_pdf_render_and_view_actual_pixels(documents):
    result = call(documents, "render_document", path="report.pdf", start_page=2, end_page=3)
    assert [f["page"] for f in result["files"]] == [2, 3]
    for file in result["files"]:
        with Image.open(documents / file["path"]) as image:
            assert image.format == "PNG" and image.width > 600
        viewed = run("view_image", {"path": file["path"]}, documents)
        assert viewed["image"].startswith("data:image/jpeg;base64,")
    viewed = run("view_image", {"path": "chart.png"}, documents)
    assert viewed["result"]["original_size"] == [2400, 800]
    assert viewed["result"]["display_size"][0] == 2048
    assert "base64" not in json.dumps(viewed["result"])


@pytest.mark.parametrize("extension", ["docx", "xlsx", "pptx"])
def test_real_office_rendering(documents, monkeypatch, extension):
    local = Path("/Applications/LibreOffice.app/Contents/MacOS")
    if local.exists():
        monkeypatch.setenv("PATH", f"{local}{os.pathsep}{os.environ['PATH']}")
    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        pytest.skip("LibreOffice is not installed")
    filename = {"docx": "contract", "xlsx": "budget", "pptx": "slides"}[extension]
    result = call(documents, "render_document", path=f"{filename}.{extension}")
    assert result["files"] and (documents / result["files"][0]["path"]).is_file()
    if extension == "docx":
        result = call(documents, path="contract.docx", start_page=1)
        assert "1200" in json.dumps(result["blocks"])
        assert result["blocks"][0]["location"]["page"] == 1


def test_real_scanned_pdf_ocr(documents):
    if not shutil.which("tesseract"):
        pytest.skip("Tesseract is not installed")
    image = Image.new("RGB", (1200, 400), "white")
    font = ImageFont.load_default(size=48)
    ImageDraw.Draw(image).text((50, 120), "APPROVED BUDGET 1200", fill="black", font=font)
    image.save(documents / "scan.pdf", "PDF")
    result = call(documents, path="scan.pdf", ocr_language="eng")
    assert result["blocks"][0]["kind"] == "ocr_text"
    assert "1200" in result["blocks"][0]["text"]
    assert result["blocks"][0]["location"]["page"] == 1


def test_path_escape_symlinks_and_bad_selectors(documents, tmp_path_factory):
    outside = tmp_path_factory.mktemp("outside") / "secret.png"
    Image.new("RGB", (10, 10)).save(outside)
    (documents / "escape.png").symlink_to(outside)
    for path in [str(outside), "escape.png", "../outside/secret.png"]:
        with pytest.raises((ValueError, FileNotFoundError)):
            call(documents, "view_image", path=path)
    (documents / "document-renders").symlink_to(outside.parent, target_is_directory=True)
    with pytest.raises(ValueError, match="Unsafe"):
        call(documents, "render_document", path="report.pdf")
    (documents / "document-renders").unlink()
    for args in [{"start_page": 0}, {"section": "Missing"}, {"sheet": "Budget"}]:
        with pytest.raises(ValueError):
            call(documents, path="contract.docx", **args)
    with pytest.raises(ValueError, match="5 pages"):
        call(documents, "render_document", path="report.pdf", end_page=6)


@pytest.fixture
def executor(documents, monkeypatch):
    hooks, recorder = MagicMock(), MagicMock()
    monkeypatch.setattr("aio_agent_platform.tools.executor.get_hook_manager", lambda: hooks)
    monkeypatch.setattr("aio_agent_platform.tools.executor.get_recorder", lambda: recorder)

    async def execute(_sandbox, command):
        argv = shlex.split(command)
        if argv[0] == "python3":
            request = json.loads(base64.b64decode(argv[-1]))
            request["root"] = str(documents)
            argv[0] = sys.executable
            argv[-1] = base64.b64encode(json.dumps(request).encode()).decode()
        else:
            # _get_file_size uses Linux stat, which is not portable to macOS.
            return SimpleNamespace(stdout="100", stderr="", exit_code=0)
        process = await asyncio.create_subprocess_exec(
            *argv, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        )
        stdout, stderr = await process.communicate()
        return SimpleNamespace(stdout=stdout.decode(), stderr=stderr.decode(), exit_code=process.returncode)

    manager = SimpleNamespace(get_or_create=AsyncMock(return_value=object()), execute=execute)
    registry = ToolRegistry()
    for tool in DOCUMENT_TOOLS:
        registry.register(tool)
    return ToolExecutor(registry, manager), hooks, recorder


async def execute_tool(executor, name, arguments, **kwargs):
    return await executor.execute(name, arguments, "tc1", "u1", "s1",
                                  workspace_id="w1", workspace_slug="project", **kwargs)


async def test_executor_real_worker_artifacts_permissions_and_image_privacy(executor):
    executor, hooks, recorder = executor
    result = await execute_tool(executor, "view_image", {"path": "chart.png"})
    assert result.success and result.image_data.startswith("data:image/jpeg;base64,")
    assert "base64" not in result.output
    assert "base64" not in repr(result)
    assert "base64" not in str(hooks.mock_calls) + str(recorder.mock_calls)
    result = await execute_tool(executor, "render_document", {"path": "report.pdf"})
    assert result.success and result.file_changes[0]["filename"] == "page-1.png"
    assert result.file_changes[0]["workspace_id"] == "w1"
    for path in ["/workspace/other/report.pdf", "../../report.pdf"]:
        result = await execute_tool(executor, "read_document", {"path": path})
        assert not result.success
    result = await execute_tool(executor, "read_document", {"path": "report.pdf"}, allowed_tools=set())
    assert not result.success and "Permission denied" in result.error
    result = await execute_tool(executor, "read_document", {"path": "missing.pdf"})
    assert not result.success and "FileNotFoundError" in result.error


@pytest.mark.parametrize("provider_type", ["openai", "anthropic"])
async def test_pixels_reach_next_model_call_but_not_history_or_events(executor, provider_type):
    executor, _, _ = executor
    provider = MagicMock()
    provider.provider_type, provider.supports_vision, provider.model = provider_type, True, "vision"
    captured = []

    async def stream(messages, tools=None):
        captured.append(list(messages))
        if len(captured) == 1:
            yield LLMChunk(type="tool_call_start", tool_call=ToolCall(id="tc1", name="view_image", arguments={}))
            yield LLMChunk(type="tool_call_delta", argument_delta='{"path":"chart.png"}')
        else:
            yield LLMChunk(type="text_delta", content="A blue chart.")
        yield LLMChunk(type="done", usage={"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2})

    provider.stream = stream
    loop = AgentLoop(provider, executor, trust_level="full", workspace_id=uuid4(), workspace_slug="project")
    history = ChatTurnRecorder(uuid4(), uuid4())
    events = []
    # A resumed run has metadata in history, never pixels: it must reread the image.
    token = completed_calls.set({call_key("view_image", {"path": "chart.png"}): {"preview": "cached metadata"}})
    try:
        async for event in loop.run("Look at chart.png", uuid4(), uuid4(), conversation_history=[], tools=[]):
            history.record(event)
            events.append(event)
    finally:
        completed_calls.reset(token)
    assert len(captured) == 2
    contents = [m.content for m in captured[1] if isinstance(m.content, list)]
    image = next(b for content in contents for b in content if b["type"] in {"image", "image_url"})
    if provider_type == "anthropic":
        pixels = image["source"]["data"]
    else:
        pixels = image["image_url"]["url"].split(",", 1)[1]
    assert Image.open(io.BytesIO(base64.b64decode(pixels))).width == 2048
    assert pixels not in str(events) + str(history.tool_calls)
    traced = sanitize_messages_for_trace([{"role": "user", "content": contents[0]}])
    assert pixels not in str(traced)


async def test_nonvision_model_fails_before_reading_pixels(executor):
    executor, _, _ = executor
    provider = MagicMock(supports_vision=False)
    loop = AgentLoop(provider, executor)
    result = await loop._execute_tool_traced(ToolCall(id="a", name="view_image", arguments={}), None)
    assert not result.success and "视觉" in result.error
    executor.sandbox_mgr.get_or_create.assert_not_awaited()
